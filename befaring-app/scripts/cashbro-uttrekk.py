#!/usr/bin/env python3
"""
cashbro-uttrekk.py — RÅ cash-kalender for HoY, sep 2026 → jun 2027. Datalag før gate/spillbrett.
Tre lag innbetalinger + kjente kostnader, hver for seg. Ingen buffer-logikk her.

  SIKKER      kjøpekontrakt signert i Oneflow (mal 5161707) → cash = overtakelsesdato + 3 arbeidsdager.
              Beløp = provisjon inkl mva fra oppgjørsarket (fallback 6 % × salgssum, min 45k).
              Allerede fakturert og betalt i PowerOffice → utelatt. Fakturert, ubetalt → kundefordring (cash = forfall, min i dag+3).
  SANNSYNLIG  aktive oppdrag (scorecard-portefølje) × P(salg i måned m) fra KM-kurven → cash måneden etter (P75 cash-lag 32 d).
  KOSTNADER   parametere (merket) + regler: meglerprovisjon ut 1. i måneden etter oppgjør, AGA terminvis, mva terminvis.

Kjør fra befaring-app/: python3 scripts/cashbro-uttrekk.py [--json ut.json]
"""
import os, sys, json, re, io, datetime as dt, urllib.request, collections
HERE = os.path.dirname(os.path.abspath(__file__)); APP = os.path.abspath(os.path.join(HERE, '..'))
ENV = {}
for line in open(os.path.join(APP, '.env')):
    line = line.strip()
    if '=' in line and not line.startswith('#'): k, v = line.split('=', 1); ENV[k.strip()] = v.strip().strip('"').strip("'")
HS_TOKEN = ENV.get('HUBSPOT_TOKEN') or open(os.path.join(APP, '..', 'HoY Internportal', 'hubspot-token.txt')).read().strip()
SB = ENV['SUPABASE_URL'].rstrip('/'); SBH = {'apikey': ENV['SUPABASE_SERVICE_ROLE_KEY'], 'Authorization': 'Bearer ' + ENV['SUPABASE_SERVICE_ROLE_KEY']}
OFH = {'x-oneflow-api-token': ENV['ONEFLOW_API_TOKEN'], 'x-oneflow-user-email': ENV['ONEFLOW_USER_EMAIL']}
DROPBOX = 'https://www.dropbox.com/scl/fi/tg66hdj0ef48nkkfaq1kf/oppgj-r-2026-l-nn-solgte-b-ter.xlsx?rlkey=oks6n4bxqqau0ofj3gu8n6m7n&dl=1'
OUT = sys.argv[sys.argv.index('--json') + 1] if '--json' in sys.argv else None
TODAY = dt.date.today(); START = TODAY.replace(day=1); END = dt.date(2027, 6, 30)

# ── PARAMETERE (til Sindre for kontroll) ─────────────────────────────────────
PAR = {
  'oppgjor_arbeidsdager': 3,                 # cash = overtakelse + N arbeidsdager
  'cash_lag_sannsynlig_dager': 32,           # P75 fra deal-for-deal
  'drift_fast_mnd': None,                    # None = beregnes fra saldobalansen (6xxx+7xxx YTD − avskrivninger) ÷ måneder hittil i år
  'fastlonn': [                              # Sindre 10. sep + arbeidsavtale (HoY-okonomi-handoff): Marte 450k fast + 10 % av Henriks provisjon, FP 10,2 %; Philip 13k, ingen feriepenger
    {'navn': 'Marte',  'brutto_mnd': 450000 / 12, 'fp_sats': 0.102, 'note': 'fast 450 000/år (arbeidsavtale) — bonus 10 % av Henriks provisjon legges på Henrik-linjene'},
    {'navn': 'Philip', 'brutto_mnd': 13000,       'fp_sats': 0.0,   'note': '13 000/mnd (Sindre) — fakturerer i tillegg per oppdrag (ligger i 4500 → direkte oppdragskost)'},
  ],
  'marte_bonus_av_henrik_provisjon': 0.10,   # 10 % av Henriks provisjon (= 4 % av Henriks omsetning), utbetales samtidig med provisjonen
  'sindre_brutto_mnd': 969498 / 12,          # 7,1 G, flatt fremover (2026-uttak hittil er lavt — «henger bak»; etterslep ikke lagt inn)
  'sindre_uttak_2026_brutto': None,          # brutto lønn Sindre har tatt ut i 2026 hittil (feriepengegrunnlag). None = ukjent → 0 til lønnsdata er synket (po_salary_lines)
  'feriepenger_2940_bruk_saldo': False,      # Sindre 10. sep: saldo 2940 er feil (gamle ansatte, Philip). False = beregn grunnlag selv: Henrik/Daniel 40 % × oms 2026, Marte, Sindre
  'lonn_dag': 1,                             # PO: lønn «September 2026» utbetalt 02.09 → lønn for mnd m ut 1. i m
  'aga_sats': 0.141, 'feriepenger_sats': 0.12,
  'feriepenger_utbetaling': '2027-06-01',    # opptjent 2026 (saldo 2940 + påløp sep–des) utbetales juni 2027 (PO 2025: «Feriepenger 2025» ut i juli)
  'billan_termin': 14781,                    # DNB EF68648: avdrag 12 922 + renter 1 859 (termin 22, aug 2026), 21. hver mnd — ikke i 6xxx/7xxx
  'leverandorgjeld_dager': [7, 30],          # saldo 2400 betales 50/50 om 7 og 30 dager
  'utbytte_avsatt': {'belop': 700000, 'dato': None, 'note': 'konto 2800 Avsatt utbytte — vedtatt, ikke utbetalt. dato=None → IKKE lagt inn i kurven, vises som varsel'},
  'megler_provisjonssats': 0.40,             # Henrik: 40 % av egen omsetning eks mva, ut 1. i mnd etter oppgjør
  'utbytte': [],                             # trekk: [{'belop':300000,'dato':'2026-11-15'}] — tomt = ingen
  'kjente_innbetalinger': [{'belop': 993540, 'dato': '2027-06-15', 'note': 'Sargo 45 Fly — nybygg, provisjon inkl mva ved levering juni 2027 (ingen kjøpekontrakt i Oneflow)'}],
  'engangs': [{'belop': 125000, 'dato': '2027-11-15', 'note': 'motfakturering Oslo Båt'}, {'belop': 50000, 'dato': '2027-10-15', 'note': 'underleverandør'}],
  'forskuddsskatt_2027': 162000,             # PO: 2025-resultat 736 978 × 22 % ≈ 162k, to terminer 15.feb/15.apr 2027 (restskatt 2024 var 194k, 2025 ga 179k til gode)
  'annullert': ['26035'],                    # kjøpekontrakt signert, men kjøper trakk seg (Sindre 9. sep) — ingen cash
  'mottatt_uten_faktura': ['Hanse 385'],     # penger mottatt, faktura glemt (Sindre) — ikke fremtidig cash
  'downside_sannsynlig_faktor': 0.77,        # P25/P50 fra scorecard-spredningen (Selskap) — Downside = sikker + 77 % av sannsynlig
  'mva_inngaende_andel_drift': 0.6,          # andel av drift som har fradragsberettiget mva
  # PLAN-lag = budsjettet/motoren, IKKE gjetting: omsetning eks mva per måned som planen krever, fylt inn der sikker+sannsynlig ikke allerede dekker.
  'plan_h2_2026': 2900000,                   # revidert H2-plan (sesongfordelt jul–des med motorens profil)
  'plan_2027_omsetning': 3800000,            # Økonomimotor «Som i dag» (Sindre 2,3M + Henrik 1,5M) — byttes til 4 981 250 ved beslutning 1+1
  'plan_henrik_andel': 1500000 / 3800000,    # andel av plan-omsetning som utløser meglerprovisjon (Henrik)
  'markedskost_per_oppdrag': 6600, 'oppdrag_per_salg': 1.3, 'inntekt_per_bat': 58375,
}
SEAS = [0.030, 0.0524, 0.0874, 0.0554, 0.1728, 0.2095, 0.1457, 0.0554, 0.0816, 0.0340, 0.0447, 0.0311]
def http(url, headers=None, body=None, method=None):
    req = urllib.request.Request(url, data=json.dumps(body).encode() if body is not None else None, headers=headers or {}, method=method or ('POST' if body is not None else 'GET'))
    with urllib.request.urlopen(req, timeout=60) as r: return r.read()
def sbq(path): return json.loads(http(f'{SB}/rest/v1/{path}', {**SBH, 'Range': '0-9999'}))
def of(path): return json.loads(http('https://api.oneflow.com/v1' + path, OFH))
def hs(path, body=None): return json.loads(http('https://api.hubapi.com' + path, {'Authorization': 'Bearer ' + HS_TOKEN, 'Content-Type': 'application/json'}, body))
def add_workdays(d, n):
    while n > 0:
        d += dt.timedelta(days=1)
        if d.weekday() < 5: n -= 1
    return d
MND = {'januar':1,'februar':2,'mars':3,'april':4,'mai':5,'juni':6,'juli':7,'august':8,'september':9,'oktober':10,'november':11,'desember':12}
def parse_no_date(s):
    if not s: return None
    m = re.match(r'(\d{1,2})\.?\s*([a-zæøå]+)\s*(\d{4})', s.strip().lower())
    if m and m.group(2) in MND: return dt.date(int(m.group(3)), MND[m.group(2)], int(m.group(1)))
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', s.strip())
    if m: return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r'(\d{1,2})[./](\d{1,2})[./](\d{4})', s.strip())
    if m: return dt.date(int(m.group(3)), int(m.group(2)), int(m.group(1)))
    return None
ym = lambda d: d.strftime('%Y-%m')
months = []; d = START
while d <= END: months.append(ym(d)); d = (d.replace(day=28) + dt.timedelta(days=4)).replace(day=1)

CAL = {m: collections.defaultdict(float) for m in months}   # CAL[mnd][linje] = beløp (+inn / −ut)
ROWS = []  # detaljer
def post(dato, linje, belop, note=''):
    if dato is None: return
    if dato < START: dato = START if dato < TODAY else dato
    k = ym(dato)
    if k not in CAL: return
    CAL[k][linje] += belop; ROWS.append({'dato': dato.isoformat(), 'linje': linje, 'belop': round(belop), 'note': note})

# ── 1. Oppgjørsark: provisjon inkl mva per oppdragsnr ────────────────────────
import openpyxl
ws = openpyxl.load_workbook(io.BytesIO(http(DROPBOX)), data_only=True).active
hdr = [c.value for c in ws[1]]; ci = {h: i for i, h in enumerate(hdr) if isinstance(h, str)}
SHEET = {}
for r in ws.iter_rows(min_row=2, values_only=True):
    if not isinstance(r[ci['Solgt dato']], dt.datetime): continue
    SHEET[str(r[ci['Oppdragsnr']] or '').strip() or ('NAVN:' + str(r[ci['Båttype']]))] = {'bat': r[ci['Båttype']], 'solgt': r[ci['Solgt dato']].date(), 'salgssum': float(r[ci['Salgssum']] or 0), 'provisjon': float(r[ci['Provisjon']] or 0), 'oms': float(r[ci['Omsetning ex.mva']] or 0), 'inn': (r[ci['Oppdrag inn']] or '').strip(), 'av': (r[ci['Solgt av']] or '').strip()}

# ── 2. PowerOffice: fakturaer (betalt / ubetalt) per prosjektkode ────────────
projs = {p['id']: p['code'] for p in sbq('po_projects?select=id,code')}
inv = sbq('po_outgoing_invoices?select=project_id,total_amount,balance,voucher_date,due_date,is_reversed&voucher_date=gte.2026-01-01')
INV = collections.defaultdict(list)
for i in inv:
    if i['is_reversed']: continue
    INV[projs.get(i['project_id'], '?')].append(i)
receivables = [i for i in inv if not i['is_reversed'] and float(i['balance'] or 0) > 0]
for i in receivables:
    due = dt.date.fromisoformat(i['due_date'] or i['voucher_date']); cash = max(due, add_workdays(TODAY, 3))
    post(cash, 'SIKKER · kundefordring (fakturert, ubetalt)', float(i['balance']), f"prosjekt {projs.get(i['project_id'],'?')} fakturert {i['voucher_date']}")
invoiced_codes = set(INV.keys())

# ── 2b. PowerOffice saldobalanse i dag (po_trial_balance, synkes 03:30) — balansepostene som blir cash ──
TB = {t['account_no']: float(t['closing_balance'] or 0) for t in sbq('po_trial_balance?select=account_no,closing_balance,as_of_date')}
TB_DATO = sbq('po_trial_balance?select=as_of_date&limit=1')[0]['as_of_date']
MND_HIA = (TODAY - dt.date(TODAY.year, 1, 1)).days / 30.44   # måneder hittil i år (for YTD-snitt)
drift_ytd = sum(v for a, v in TB.items() if 6000 <= a < 8000 and a not in (6010, 6011, 6012, 6013))
direkte_netto_ytd = TB.get(4500, 0) + TB.get(4300, 0) + TB.get(4060, 0) + TB.get(4360, 0) + TB.get(3000, 0)   # fremmedytelser minus viderefakturert (3000 er negativ)
personal_annet_ytd = sum(TB.get(a, 0) for a in (5500, 5910, 5920, 5950, 5990))
DRIFT_MND = PAR['drift_fast_mnd'] or round(drift_ytd / MND_HIA)
DIREKTE_MND = round(max(0, direkte_netto_ytd) / MND_HIA); PERSONAL_ANNET_MND = round(personal_annet_ytd / MND_HIA)
# hovedbok september (for å skille termin 4 fra termin 5 i mva-posisjonen, og termin 4 fra sept-AGA)
T0 = START.month if START.month % 2 == 1 else START.month - 1          # første måned i inneværende mva/AGA-termin
T_START = dt.date(START.year, T0, 1)
hb_sep = sbq(f"po_account_transactions?select=account_no,amount&is_reversed=eq.false&posting_date=gte.{T_START.isoformat()}")
sep_mva = sum(r['amount'] for r in hb_sep if r['account_no'] in (2701, 2702, 2711, 2712, 2714))
sep_aga = sum(r['amount'] for r in hb_sep if r['account_no'] == 5400)
MVA_POS = sum(TB.get(a, 0) for a in (2701, 2702, 2711, 2712, 2714, 2740))   # negativ = skyldig
AGA_ACC = collections.defaultdict(float)   # AGA påløpt per mnd → betales termin
FP_ACC = 0.0                               # feriepenger påløpt på lønn utbetalt resten av 2026 → juni 2027
def lonn(dato, linje, brutto, note='', fp_sats=None):
    """Brutto lønn ut på dato (netto + skattetrekk samme dag). AGA 14,1 % til termin, feriepenger til juni-klumpen (kun 2026-opptjening)."""
    global FP_ACC
    if dato < TODAY: return
    post(dato, linje, -brutto, note); AGA_ACC[ym(dato)] += brutto * PAR['aga_sats']
    if dato.year == 2026: FP_ACC += brutto * (PAR['feriepenger_sats'] if fp_sats is None else fp_sats)
def henrik_prov(dato, linje, brutto, note=''):
    """Henriks provisjon + Martes bonus (10 % av provisjonen), begge 1. i måneden etter oppgjør."""
    lonn(dato, linje, brutto, note)
    lonn(dato, 'UT · Marte bonus (10 % av Henriks provisjon)', brutto * PAR['marte_bonus_av_henrik_provisjon'], note, fp_sats=0.102)
# feriepengegrunnlag 2026 hittil (Sindre: saldo 2940 er feil — kun Henrik, Sindre, Marte, Daniel har krav)
def prov_base(navn):
    b = 0.0
    for nr, s in SHEET.items():
        if s['solgt'].year != 2026: continue
        inn, av = s['inn'], s['av']
        if av == navn: b += s['oms'] * (1.0 if inn == av or (inn == 'Daniel' and s['solgt'].isoformat() >= '2026-08-16') else 0.5)
        elif inn == navn and av != navn and not (navn == 'Daniel' and s['solgt'].isoformat() >= '2026-08-16'): b += s['oms'] * 0.5
    return b
FP_HITTIL = {
    'Henrik': prov_base('Henrik') * PAR['megler_provisjonssats'] * PAR['feriepenger_sats'],
    'Daniel': prov_base('Daniel') * PAR['megler_provisjonssats'] * PAR['feriepenger_sats'],
    'Marte':  (prov_base('Henrik') * PAR['megler_provisjonssats'] * PAR['marte_bonus_av_henrik_provisjon'] + 450000 / 12 * max(0, (TODAY.year - 2026) * 12 + TODAY.month - 5)) * 0.102,   # fast fra mai + bonus
    'Sindre': (PAR['sindre_uttak_2026_brutto'] or 0) * PAR['feriepenger_sats'],
}

# ── 3. Oneflow: kjøpekontrakter signert 2026 → overtakelsesdato + salgssum ───
first = of('/contracts?limit=100&offset=0'); total = first['count']; contracts = list(first['data'])
for o in range(100, total, 100): contracts += of(f'/contracts?limit=100&offset={o}')['data']
tid = lambda c: int((c.get('_private_ownerside') or {}).get('template_id') or 0)
nm = lambda c: (c.get('_private') or {}).get('name') or c.get('name') or ''
KK = [c for c in contracts if tid(c) == 5161707 and c.get('state') == 'signed' and (c.get('state_updated_time') or '') >= '2026-01-01']
OP = {}  # overtakelsesprotokoll signert per nr
for c in contracts:
    if tid(c) == 5137684 and c.get('state') == 'signed':
        m = re.match(r'^(\d{5})', nm(c).strip());
        if m: OP[m.group(1)] = c['state_updated_time'][:10]
SIKKER = []
for c in KK:
    df = of(f'/contracts/{c["id"]}/data_fields').get('data', [])
    f = {(x.get('custom_id') or x.get('name')): (x.get('value') or '') for x in df}
    overt = parse_no_date(f.get('Deal_Dato for overtakelse')); salgssum = float(re.sub(r'\D', '', f.get('Deal_Salgssum') or '0') or 0)
    m = re.match(r'^(\d{5})', nm(c).strip()); nr = m.group(1) if m else None
    if nr in PAR['annullert'] or any(k.lower() in nm(c).lower() for k in PAR['mottatt_uten_faktura']):
        SIKKER.append({'nr': nr, 'navn': nm(c), 'status': 'annullert' if nr in PAR['annullert'] else 'mottatt uten faktura', 'overtakelse': None, 'provisjon_inkl': 0, 'solgt_av': None}); continue
    sh = SHEET.get(nr) if nr else None
    prov = sh['provisjon'] if sh and sh['provisjon'] else max(45000, salgssum * 0.06)
    oms = sh['oms'] if sh and sh['oms'] else prov / 1.25
    paid = nr in invoiced_codes and all(float(i['balance'] or 0) == 0 for i in INV[nr])
    invoiced_open = nr in invoiced_codes and not paid
    row = {'nr': nr, 'navn': nm(c), 'signert_kk': c['state_updated_time'][:10], 'overtakelse': overt.isoformat() if overt else None, 'salgssum': salgssum, 'provisjon_inkl': round(prov), 'oms_eks': round(oms), 'i_arket': bool(sh), 'solgt_av': sh['av'] if sh else None, 'fakturert': nr in invoiced_codes, 'betalt': paid, 'levert': OP.get(nr)}
    if paid: row['status'] = 'betalt'; SIKKER.append(row); continue
    if invoiced_open: row['status'] = 'kundefordring (i pkt 2)'; SIKKER.append(row); continue
    if not overt: row['status'] = 'MANGLER overtakelsesdato'; SIKKER.append(row); continue
    cash = add_workdays(overt, PAR['oppgjor_arbeidsdager'])
    if cash < TODAY: cash = add_workdays(TODAY, 3); row['status'] = 'overtatt, ikke fakturert ennå → antatt cash om 3 ad'
    else: row['status'] = 'venter overtakelse'
    row['cash'] = cash.isoformat(); SIKKER.append(row)
    post(cash, 'SIKKER · provisjon ved overtakelse', prov, f"{nr} {row['navn'][:30]} overt. {overt}")
    # meglerprovisjon ut 1. i måneden etter oppgjør (Henrik-solgte); Sindre = 7,1 G flatt
    if (row['solgt_av'] or '').lower().startswith('henrik'):
        pay = (cash.replace(day=1) + dt.timedelta(days=32)).replace(day=1)
        henrik_prov(pay, 'UT · meglerprovisjon (Henrik 40 %)', oms * PAR['megler_provisjonssats'], f"{nr}")
    # utgående mva på provisjonen → termin
    post(cash, '_mva_ut', -(prov - oms))

# ── 4. Sannsynlig: portefølje × P(salg i mnd) fra KM (scorecard_state) ───────
st = sbq('scorecard_state?select=state&id=eq.1')[0]['state']
kurver = st.get('kurver', {}).get('klasser', {})
# rekonstruer grov månedlig hasard fra 90/180/365-punktene (lineær mellom punkter, flat etter 365)
def S_at(k, t):
    c = kurver.get(k) or kurver.get('ukjent'); pts = [(0, 1.0), (90, 1 - c['solgt_innen_90']), (180, 1 - c['solgt_innen_180']), (365, 1 - c['solgt_innen_365'])]
    if t >= 365: return pts[-1][1]
    for i in range(1, len(pts)):
        if t <= pts[i][0]: (t0, s0), (t1, s1) = pts[i - 1], pts[i]; return s0 + (s1 - s0) * (t - t0) / (t1 - t0)
    return pts[-1][1]
SANN = []
for o in st['portefolje']['oppdrag']:
    if not o.get('pris'): continue
    prov = max(45000, o['pris'] * 0.06); oms = prov / 1.25; alder0 = o['alder_dager']; s0 = S_at(o['prisklasse'], alder0)
    if s0 <= 0: continue
    for m in months:
        mstart = dt.date.fromisoformat(m + '-01'); mend = (mstart.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)
        t1, t2 = alder0 + (mstart - TODAY).days, alder0 + (mend - TODAY).days
        p = max(0.0, (S_at(o['prisklasse'], max(alder0, t1)) - S_at(o['prisklasse'], t2)) / s0)
        if p <= 0: continue
        cash = mstart + dt.timedelta(days=15 + PAR['cash_lag_sannsynlig_dager'])
        post(cash, 'SANNSYNLIG · portefølje × P', prov * p, f"{o['nr']} p={p:.2f}")
        if o['megler'] == 'Henrik':
            pay = (cash.replace(day=1) + dt.timedelta(days=32)).replace(day=1)
            henrik_prov(pay, 'UT · meglerprovisjon sannsynlig (Henrik)', oms * p * PAR['megler_provisjonssats'], o['nr'])
        post(cash, '_mva_ut_sann', -(prov - oms) * p)
    SANN.append({'nr': o['nr'], 'navn': o['navn'], 'megler': o['megler'], 'pris': o['pris'], 'p_i_ar': o['p_salg_i_ar']})

# ── 4b. PLAN-lag: planens omsetning per måned minus det sikker+sannsynlig allerede dekker ──
h2w = sum(SEAS[6:])
plan_oms = {}
for m in months:
    y, mo = int(m[:4]), int(m[5:])
    plan_oms[m] = PAR['plan_h2_2026'] * SEAS[mo - 1] / h2w if y == 2026 else PAR['plan_2027_omsetning'] * SEAS[mo - 1]
PLAN = []
for m in months:
    # Dekket = det motoren allerede har i måned m: sikker-salg cashes samme måned (overtakelse + 3 ad),
    # sannsynlig-salg i m cashes i m+1 (32 d). Begge omregnes til eks mva.
    nxt = ym((dt.date.fromisoformat(m + '-01').replace(day=28) + dt.timedelta(days=4)).replace(day=1))
    dekket_inkl = CAL[m].get('SIKKER · provisjon ved overtakelse', 0) + (CAL[nxt].get('SANNSYNLIG · portefølje × P', 0) if nxt in CAL else 0)
    fyll = max(0.0, plan_oms[m] - dekket_inkl / 1.25)
    PLAN.append({'mnd': m, 'plan_oms_eks': round(plan_oms[m]), 'dekket_eks': round(dekket_inkl / 1.25), 'plan_fyll_eks': round(fyll)})
    if fyll <= 0 or nxt not in CAL: continue
    cash = dt.date.fromisoformat(nxt + '-01') + dt.timedelta(days=15)
    post(cash, 'PLAN · nye oppdrag (motor − dekket)', fyll * 1.25, f'{m}: plan {plan_oms[m]:,.0f} − dekket {dekket_inkl/1.25:,.0f}')
    post(cash, '_mva_ut_plan', -fyll * 0.25)
    pay = (cash.replace(day=1) + dt.timedelta(days=32)).replace(day=1)
    henrik_prov(pay, 'UT · meglerprovisjon plan (Henrik-andel)', fyll * PAR['plan_henrik_andel'] * PAR['megler_provisjonssats'], m)
    post(dt.date.fromisoformat(m + '-01'), 'UT · markedskost nye oppdrag (plan)', -fyll / PAR['inntekt_per_bat'] * PAR['oppdrag_per_salg'] * PAR['markedskost_per_oppdrag'], m)

# ── 5. Kjente kostnader ───────────────────────────────────────────────────────
def ut(dato, linje, belop, note=''):
    """Kostnad ut — hopper over datoer før i dag (allerede betalt, ligger i live banksaldo)."""
    if dato >= TODAY: post(dato, linje, -belop, note)
dim = lambda d: ((d.replace(day=28) + dt.timedelta(days=4)).replace(day=1) - dt.timedelta(days=1)).day
for m in months:
    d1 = dt.date.fromisoformat(m + '-01')
    # drift: YTD-snitt fra saldobalansen, fordelt jevnt over måneden → i inneværende måned kun gjenstående dager
    andel = (dim(d1) - TODAY.day + 1) / dim(d1) if m == ym(TODAY) else 1.0
    # regnskapstall er eks mva → cash ut = netto + 25 % mva på den mva-pliktige andelen; mva-en kommer tilbake i terminoppgjøret (_mva_inn)
    mva_drift = DRIFT_MND * andel * PAR['mva_inngaende_andel_drift'] * 0.25
    post(max(d1, TODAY), 'UT · drift (6xxx+7xxx YTD-snitt, eks avskrivn., inkl mva)', -(DRIFT_MND * andel + mva_drift), f'saldobalanse {TB_DATO}: {drift_ytd:,.0f} ÷ {MND_HIA:.1f} mnd')
    post(max(d1, TODAY), 'UT · direkte oppdragskost netto (4xxx − viderefakt.)', -DIREKTE_MND * andel, 'saldobalanse YTD')
    post(max(d1, TODAY), 'UT · personalkost annet (OTP, forsikr., kantine)', -PERSONAL_ANNET_MND * andel, 'saldobalanse YTD 55xx/59xx')
    post(max(d1, TODAY), '_mva_inn', mva_drift)
    ut(d1.replace(day=21), 'UT · billån DNB (avdrag + renter)', PAR['billan_termin'], 'konto 2242/8151, termin 21. hver mnd')
    # lønn 1. i måneden: Sindre + faste
    ld = d1.replace(day=PAR['lonn_dag'])
    lonn(ld, 'UT · Sindre lønn 7,1 G', PAR['sindre_brutto_mnd'], 'param')
    for f in PAR['fastlonn']: lonn(ld, f'UT · fastlønn {f["navn"]}', f['brutto_mnd'], f['note'], fp_sats=f['fp_sats'])
# AGA: skyldig i dag (2770) = forrige termin (ubetalt til 15. i termin-start-mnd) + påløpt i inneværende termin
aga_prev = -TB.get(2770, 0) - sep_aga
prev_due = dt.date(START.year, T0, 15)
if aga_prev > 0 and prev_due >= TODAY: ut(prev_due, 'UT · AGA termin', aga_prev, f'forrige termin fra saldo 2770 {TB.get(2770,0):,.0f} − påløpt i inneværende termin {sep_aga:,.0f}')
AGA_ACC[ym(T_START)] += sep_aga
for (y, mo) in sorted({(int(m[:4]), int(m[5:])) for m in months}):
    if mo % 2 == 0:
        v = AGA_ACC.get(f'{y}-{mo-1:02d}', 0) + AGA_ACC.get(f'{y}-{mo:02d}', 0)
        pm = mo + 1; py = y + (1 if pm > 12 else 0); pm = 1 if pm > 12 else pm
        if v > 0: ut(dt.date(py, pm, 15), 'UT · AGA termin', v, f'termin {mo//2}/{y} (lønn {mo-1}–{mo})')
# feriepenger: opptjent 2026 = saldo 2940 i dag + 12 % av lønn utbetalt resten av året → ut juni 2027 (AGA på feriepenger til termin 3/2027 = 15.7, utenfor horisont)
fp_hittil = -TB.get(2940, 0) if PAR['feriepenger_2940_bruk_saldo'] else sum(FP_HITTIL.values())
fp = fp_hittil + FP_ACC
ut(dt.date.fromisoformat(PAR['feriepenger_utbetaling']), 'UT · feriepenger (opptjent 2026)', fp, ('saldo 2940' if PAR['feriepenger_2940_bruk_saldo'] else 'beregnet: ' + ', '.join(f'{k} {v:,.0f}' for k, v in FP_HITTIL.items())) + f' + påløp resten av året {FP_ACC:,.0f}')
# leverandørgjeld i dag (2400) → 50/50
lg = -TB.get(2400, 0)
for dd in PAR['leverandorgjeld_dager']: ut(TODAY + dt.timedelta(days=dd), 'UT · leverandørgjeld (saldo 2400)', lg / len(PAR['leverandorgjeld_dager']), f'saldo {lg:,.0f}')
for e in PAR['engangs']: ut(dt.date.fromisoformat(e['dato']), 'UT · engangspost', e['belop'], e['note'])
for e in PAR['kjente_innbetalinger']: post(dt.date.fromisoformat(e['dato']), 'SIKKER · kjent innbetaling (param)', e['belop'], e['note']); post(dt.date.fromisoformat(e['dato']), '_mva_ut', -e['belop'] * 0.2)
for u in PAR['utbytte']: ut(dt.date.fromisoformat(u['dato']), 'UT · utbytte (trekk)', u['belop'], 'trekk')
if PAR['utbytte_avsatt']['dato']: ut(dt.date.fromisoformat(PAR['utbytte_avsatt']['dato']), 'UT · utbytte avsatt (konto 2800)', PAR['utbytte_avsatt']['belop'], PAR['utbytte_avsatt']['note'])
if PAR['forskuddsskatt_2027']: ut(dt.date(2027, 2, 15), 'UT · forskuddsskatt (ANSLAG)', PAR['forskuddsskatt_2027'] / 2); ut(dt.date(2027, 4, 15), 'UT · forskuddsskatt (ANSLAG)', PAR['forskuddsskatt_2027'] / 2)
# mva-oppgjør terminvis: termin (jan-feb → 10. apr, mar-apr → 10. jun, mai-jun → 31. aug, jul-aug → 10. okt, sep-okt → 10. des, nov-des → 10. feb)
TERM = {1: (4, 10), 3: (6, 10), 5: (8, 31), 7: (10, 10), 9: (12, 10), 11: (2, 10)}
netto = collections.defaultdict(float)
# åpningsposisjon fra saldobalansen: alt bokført t.o.m. i dag. Det som er bokført i inneværende termin hører dit; resten er forrige termin (ubetalt)
prev_t0 = T0 - 2 if T0 > 1 else 11; prev_y = START.year if T0 > 1 else START.year - 1
netto[(prev_y, prev_t0)] += MVA_POS - sep_mva; netto[(START.year, T0)] += sep_mva
for m in months:
    y, mo = int(m[:4]), int(m[5:]); t0 = mo if mo % 2 == 1 else mo - 1
    netto[(y, t0)] += CAL[m].get('_mva_ut', 0) + CAL[m].get('_mva_ut_sann', 0) + CAL[m].get('_mva_ut_plan', 0) + CAL[m].get('_mva_inn', 0)
for (y, t0), v in netto.items():
    pm, pd = TERM[t0]; py = y + (1 if pm < t0 else 0)
    if dt.date(py, pm, pd) >= TODAY: post(dt.date(py, pm, pd), 'UT · mva-oppgjør (alle lag)', v, f'termin {t0}-{t0+1}/{y}')
for m in months:
    for k in ['_mva_ut', '_mva_ut_sann', '_mva_ut_plan', '_mva_inn']: CAL[m].pop(k, None)

# ── 6. Bank og saldo-kurve ────────────────────────────────────────────────────
bank = st.get('likviditet', {}).get('reell_bank'); kilde = st.get('likviditet', {}).get('bank_kilde')
LINES = sorted({k for m in months for k in CAL[m]})
SIK = [k for k in LINES if k.startswith('SIKKER')]; SAN = [k for k in LINES if k.startswith('SANNSYNLIG') or k.startswith('UT · meglerprovisjon sannsynlig')]
PL = [k for k in LINES if k.startswith('PLAN') or k.startswith('UT · meglerprovisjon plan') or k.startswith('UT · markedskost nye')]
UT = [k for k in LINES if k.startswith('UT') and k not in SAN and k not in PL]
run_d, run_b, run_p = bank, bank, bank; curve = []
for m in months:
    sik = sum(CAL[m][k] for k in SIK); san = sum(CAL[m][k] for k in SAN); pl = sum(CAL[m][k] for k in PL); ut = sum(CAL[m][k] for k in UT)
    run_d += sik + san * PAR['downside_sannsynlig_faktor'] + ut; run_b += sik + san + ut; run_p += sik + san + pl + ut
    curve.append({'mnd': m, 'sikker_inn': round(sik), 'sannsynlig_netto': round(san), 'plan_netto': round(pl), 'kostnader': round(ut), 'saldo_downside': round(run_d), 'saldo_base': round(run_b), 'saldo_plan': round(run_p)})

print(f'Bank i dag ({kilde}): {bank:,.0f}   · bokført 1920 i saldobalansen {TB_DATO}: {TB.get(1920,0):,.0f}')
print(f'SALDOBALANSE → cash: kundefordringer 1500 {TB.get(1500,0):,.0f} · leverandørgjeld 2400 {TB.get(2400,0):,.0f} · skyldig AGA 2770 {TB.get(2770,0):,.0f} · mva-posisjon 27xx {MVA_POS:,.0f} (herav inneværende termin {sep_mva:,.0f})')
print(f'  feriepenger 2026 beregnet hittil: ' + ', '.join(f'{k} {v:,.0f}' for k, v in FP_HITTIL.items()) + f' = {sum(FP_HITTIL.values()):,.0f} (saldo 2940 sier {-TB.get(2940,0):,.0f} — Sindre: feil, inkluderer gamle ansatte/Philip)')
print(f'  skyldige feriepenger 2940 {TB.get(2940,0):,.0f} (+AGA 2785 {TB.get(2785,0):,.0f}) · avsatt utbytte 2800 {TB.get(2800,0):,.0f} · annen kortsiktig gjeld 2990 {TB.get(2990,0):,.0f} · billån 2242 {TB.get(2242,0):,.0f} · betalbar skatt 2500 {TB.get(2500,0):,.0f}')
print(f'  drift 6xxx+7xxx YTD {drift_ytd:,.0f} → {DRIFT_MND:,.0f}/mnd · direkte netto {DIREKTE_MND:,.0f}/mnd · personal annet {PERSONAL_ANNET_MND:,.0f}/mnd · lønn 5000 YTD {TB.get(5000,0):,.0f}')
if not PAR['utbytte_avsatt']['dato']: print(f'  ⚠ AVSATT UTBYTTE {PAR["utbytte_avsatt"]["belop"]:,.0f} (konto 2800) er IKKE lagt inn — sett dato i PAR["utbytte_avsatt"] når det skal betales')
if TB.get(2990, 0): print(f'  ⚠ ANNEN KORTSIKTIG GJELD 2990 {TB.get(2990,0):,.0f} — ikke lagt inn (hva er det?)')
print(f'parametere: {json.dumps({k:v for k,v in PAR.items() if k not in ("engangs","utbytte","fastlonn")}, ensure_ascii=False)}\n')
print(f'{"mnd":8}{"sikker":>10}{"sannsynl":>10}{"plan":>10}{"kost":>10}{"DOWNSIDE":>11}{"BASE":>11}{"BASE+PLAN":>11}')
for c in curve: print(f'{c["mnd"]:8}{c["sikker_inn"]:>10,}{c["sannsynlig_netto"]:>10,}{c["plan_netto"]:>10,}{c["kostnader"]:>10,}{c["saldo_downside"]:>11,}{c["saldo_base"]:>11,}{c["saldo_plan"]:>11,}')
print('\nKOSTNADSLINJER per måned:'); print(f'{"mnd":8}' + ''.join(f'{k.replace("UT · ","")[:22]:>24}' for k in UT + PL))
for m in months: print(f'{m:8}' + ''.join(f'{CAL[m].get(k,0):>24,.0f}' for k in UT + PL))
print('\nPLAN-lag (eks mva): ' + ' | '.join(f"{p['mnd'][2:]}: plan {p['plan_oms_eks']/1000:.0f}k − dekket {p['dekket_eks']/1000:.0f}k = fyll {p['plan_fyll_eks']/1000:.0f}k" for p in PLAN))
lo_d = min(curve, key=lambda c: c['saldo_downside']); lo_b = min(curve, key=lambda c: c['saldo_base']); lo_p = min(curve, key=lambda c: c['saldo_plan'])
print(f'\nLAVESTE — Downside (sikker + 77 % sannsynlig − kost): {lo_d["saldo_downside"]:,.0f} i {lo_d["mnd"]} · Base (+ sannsynlig): {lo_b["saldo_base"]:,.0f} i {lo_b["mnd"]} · Base+Plan (motor fyller opp): {lo_p["saldo_plan"]:,.0f} i {lo_p["mnd"]}')
print('\nSIKKER — kjøpekontrakter 2026:')
for r in sorted(SIKKER, key=lambda x: x.get('overtakelse') or '9'):
    print(f"  {r['nr'] or '—':6} {r['navn'][:38]:38} overt={r['overtakelse'] or '?':10} cash={r.get('cash','—'):10} prov={r['provisjon_inkl']:>9,.0f} {r['solgt_av'] or '':7} {r['status']}")
print(f'\nKundefordringer (ubetalte fakturaer): {len(receivables)} stk, {sum(float(i["balance"]) for i in receivables):,.0f}')
if OUT:
    json.dump({'generert': TODAY.isoformat(), 'bank': bank, 'bank_kilde': kilde, 'param': PAR, 'kalender': {m: dict(CAL[m]) for m in months}, 'kurve': curve, 'sikker': SIKKER, 'sannsynlig': SANN, 'plan': PLAN, 'rader': ROWS}, open(OUT, 'w'), ensure_ascii=False, indent=1, default=str)
    print('skrevet', OUT)
