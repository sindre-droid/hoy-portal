#!/usr/bin/env python3
"""
scorecard-uttrekk.py — RÅ uketall til HoY-scorecardet, per ISO-uke og megler.
Kun uttrekk + telling etter LÅSTE definisjoner (8. sep 2026). Ingen mål, ingen farger.

Kolonner (kilde):
  kontakter      «call» i HubSpot, hs_timestamp i uken, per hubspot_owner_id (alle + FSBO Outreach)
  leads          deals opprettet i Pipeline A (3205247197), createdate i uken, per owner
  tilbud_sendt   Oneflow oppdragsavtale (template 5130587) published_time i uken, megler = HoY-deltaker
  signert        Oneflow OA, oppdragsgivers signatur (siste ikke-@h-y.no sign_state_updated_time) i uken
  solgt          oppgjørsark (Dropbox), Solgt dato i uken, per «Solgt av»
  provisjon      oppgjørsark, Omsetning ex.mva, KREDITERT (50/50 når Oppdrag inn ≠ Solgt av)
  publisert_7d   OA signert i uken → FINN/HubSpot publiseringsdato ≤ 7 dg (best effort via oppdrag_livslop)

Kjør fra befaring-app/: python3 scripts/scorecard-uttrekk.py [--fra 2026-07-01] [--json ut.json]
"""
import os, sys, json, re, datetime as dt, urllib.request, urllib.parse, io, ssl
from collections import defaultdict

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.abspath(os.path.join(HERE, '..'))
ENV = {}
for line in open(os.path.join(APP, '.env')):
    line = line.strip()
    if '=' in line and not line.startswith('#'):
        k, v = line.split('=', 1); ENV[k.strip()] = v.strip().strip('"').strip("'")
HS_TOKEN = ENV.get('HUBSPOT_TOKEN') or open(os.path.join(APP, '..', 'HoY Internportal', 'hubspot-token.txt')).read().strip()

DROPBOX = 'https://www.dropbox.com/scl/fi/tg66hdj0ef48nkkfaq1kf/oppgj-r-2026-l-nn-solgte-b-ter.xlsx?rlkey=oks6n4bxqqau0ofj3gu8n6m7n&dl=1'
PIPELINE_A = '3205247197'
OA_TEMPLATE = 5130587
OWNERS = {'633479117': 'Sindre', '77221549': 'Henrik', '33931214': 'Marte', '29136352': 'Daniel', '78018793': 'Philip'}
EMAILS = {'sindre@h-y.no': 'Sindre', 'henrik@h-y.no': 'Henrik', 'marte@h-y.no': 'Marte', 'daniel@h-y.no': 'Daniel', 'philip@h-y.no': 'Philip'}
# Henrik+Marte er én enhet i planen
UNIT = {'Sindre': 'Sindre', 'Henrik': 'Henrik', 'Marte': 'Henrik', 'Daniel': 'Daniel', 'Philip': 'Philip'}

args = sys.argv[1:]
FRA = dt.date.fromisoformat(args[args.index('--fra') + 1]) if '--fra' in args else dt.date(2026, 7, 1)
OUT = args[args.index('--json') + 1] if '--json' in args else None
TODAY = dt.date.today()

def isoweek(d):
    if isinstance(d, dt.datetime): d = d.date()
    y, w, _ = d.isocalendar(); return f'{y}-U{w:02d}'
def parse_ts(s):
    if not s: return None
    s = s.replace('Z', '+00:00')
    try: return dt.datetime.fromisoformat(s)
    except Exception: return None
def http(url, headers=None, body=None, method=None):
    data = json.dumps(body).encode() if body is not None else None
    req = urllib.request.Request(url, data=data, headers=headers or {}, method=method or ('POST' if data else 'GET'))
    with urllib.request.urlopen(req, timeout=60) as r: return r.read()
def hs(path, body=None):
    H = {'Authorization': 'Bearer ' + HS_TOKEN, 'Content-Type': 'application/json'}
    return json.loads(http('https://api.hubapi.com' + path, H, body))
def hs_search_all(obj, filters, props):
    out, after = [], None
    for _ in range(40):
        body = {'filterGroups': [{'filters': filters}], 'properties': props, 'limit': 100}
        if after: body['after'] = after
        j = hs(f'/crm/v3/objects/{obj}/search', body)
        out += j.get('results', [])
        after = j.get('paging', {}).get('next', {}).get('after')
        if not after: break
    return out

W = defaultdict(lambda: defaultdict(lambda: defaultdict(float)))  # W[uke][megler][kol]
NOTES = []
def add(week, who, col, n=1): W[week][who][col] += n

# ── 1. Oppgjørsark (Dropbox) ────────────────────────────────────────────────
import openpyxl
xl = http(DROPBOX)
ws = openpyxl.load_workbook(io.BytesIO(xl), data_only=True).active
hdr = [c.value for c in ws[1]]
ci = {h: i for i, h in enumerate(hdr) if isinstance(h, str)}
sales = []
for r in ws.iter_rows(min_row=2, values_only=True):
    d = r[ci['Solgt dato']]
    if not isinstance(d, dt.datetime): continue
    inn, av = (r[ci['Oppdrag inn']] or '').strip(), (r[ci['Solgt av']] or '').strip()
    oms = float(r[ci['Omsetning ex.mva']] or 0)
    sales.append({'nr': r[ci['Oppdragsnr']], 'bat': r[ci['Båttype']], 'dato': d.date().isoformat(), 'inn': inn, 'av': av, 'salgssum': r[ci['Salgssum']], 'oms': oms})
    if d.date() < FRA: continue
    wk = isoweek(d)
    add(wk, UNIT.get(av, av or 'ukjent'), 'solgt', 1)
    if inn and av and inn != av:
        add(wk, UNIT.get(inn, inn), 'provisjon', oms / 2); add(wk, UNIT.get(av, av), 'provisjon', oms / 2)
    else:
        add(wk, UNIT.get(av, av or 'ukjent'), 'provisjon', oms)
NOTES.append(f'Oppgjørsark: {len(sales)} salg i arket, siste solgt {max(s["dato"] for s in sales)}; YTD oms ex mva {sum(s["oms"] for s in sales):,.0f}')

# ── 2. HubSpot: leads (A-deals opprettet) + kontakter (calls) ───────────────
fra_iso = dt.datetime.combine(FRA, dt.time()).isoformat() + 'Z'
deals = hs_search_all('deals', [{'propertyName': 'pipeline', 'operator': 'EQ', 'value': PIPELINE_A},
                                {'propertyName': 'createdate', 'operator': 'GTE', 'value': fra_iso}],
                      ['createdate', 'hubspot_owner_id', 'dealname', 'dealstage'])
for d in deals:
    p = d['properties']; t = parse_ts(p.get('createdate'))
    add(isoweek(t), UNIT.get(OWNERS.get(p.get('hubspot_owner_id'), 'ukjent'), 'ukjent'), 'leads')
NOTES.append(f'HubSpot Pipeline A: {len(deals)} deals opprettet siden {FRA}')

calls = hs_search_all('calls', [{'propertyName': 'hs_timestamp', 'operator': 'GTE', 'value': fra_iso}],
                      ['hs_timestamp', 'hubspot_owner_id', 'hs_activity_type', 'hs_call_status', 'hs_call_direction'])
for c in calls:
    p = c['properties']; t = parse_ts(p.get('hs_timestamp'))
    who = UNIT.get(OWNERS.get(p.get('hubspot_owner_id'), 'ukjent'), 'ukjent')
    add(isoweek(t), who, 'kontakter')
    if (p.get('hs_activity_type') or '').lower().startswith('fsbo'): add(isoweek(t), who, 'kontakter_fsbo')
NOTES.append(f'HubSpot calls: {len(calls)} siden {FRA} (aktivitetstyper: {sorted(set(str(c["properties"].get("hs_activity_type")) for c in calls))})')

# ── 3. Oneflow: tilbud sendt + signert (oppdragsavtale) ─────────────────────
OFH = {'x-oneflow-api-token': ENV['ONEFLOW_API_TOKEN'], 'x-oneflow-user-email': ENV['ONEFLOW_USER_EMAIL']}
def of(path): return json.loads(http('https://api.oneflow.com/v1' + path, OFH))
first = of('/contracts?limit=100&offset=0'); total = first.get('count', 0)
contracts = list(first.get('data', []))
off = 100
while off < total and off < 1500:   # nyeste først? Oneflow sorterer ikke garantert — vi henter alt (≈11 sider)
    contracts += of(f'/contracts?limit=100&offset={off}').get('data', []); off += 100
def tid(c): return int((c.get('_private_ownerside') or {}).get('template_id') or 0)
def cname(c): return (c.get('_private') or {}).get('name') or c.get('name') or ''
def hoy_broker(c):
    for party in c.get('parties') or []:
        for pt in party.get('participants') or []:
            e = (pt.get('email') or '').lower()
            if e in EMAILS: return UNIT[EMAILS[e]]
    return 'ukjent'
oas = [c for c in contracts if tid(c) == OA_TEMPLATE]
sent_in_window = [c for c in oas if parse_ts(c.get('published_time')) and parse_ts(c['published_time']).date() >= FRA]
for c in sent_in_window:
    add(isoweek(parse_ts(c['published_time'])), hoy_broker(c), 'tilbud_sendt')
signed_recent = [c for c in oas if c.get('state') == 'signed' and parse_ts(c.get('state_updated_time')) and parse_ts(c['state_updated_time']).date() >= FRA - dt.timedelta(days=14)]
signed_rows = []
for c in signed_recent:
    det = of(f'/contracts/{c["id"]}')
    times = []
    for party in det.get('parties') or []:
        for pt in party.get('participants') or []:
            if not pt.get('signatory') or pt.get('sign_state') != 'signed': continue
            if (pt.get('email') or '').lower().endswith('@h-y.no'): continue
            if pt.get('sign_state_updated_time'): times.append(pt['sign_state_updated_time'])
    ts = parse_ts(sorted(times)[-1]) if times else parse_ts(c['state_updated_time'])
    if ts.date() < FRA: continue
    who = hoy_broker(det)
    add(isoweek(ts), who, 'signert')
    nr = (re.match(r'^(\d{5})', cname(c).strip()) or [None, None])[1]
    signed_rows.append({'id': c['id'], 'navn': cname(c), 'nr': nr, 'signert': ts.date().isoformat(), 'megler': who, 'kilde_sign': 'oppdragsgiver' if times else 'state_updated_time'})
NOTES.append(f'Oneflow: {len(contracts)}/{total} kontrakter lest, {len(oas)} oppdragsavtaler; {len(sent_in_window)} sendt siden {FRA}, {len(signed_rows)} signert siden {FRA}')

# ── 4. Publisert ≤7 dg — best effort via oppdrag_livslop (FINN-fasit) + boats.finn_kode ──
SB = ENV['SUPABASE_URL'].rstrip('/'); SBH = {'apikey': ENV['SUPABASE_SERVICE_ROLE_KEY'], 'Authorization': 'Bearer ' + ENV['SUPABASE_SERVICE_ROLE_KEY']}
nrs = [r['nr'] for r in signed_rows if r['nr']]
liv = {}
if nrs:
    q = f'{SB}/rest/v1/oppdrag_livslop?select=oppdragsnr,oppdragsavtale_signert,annonse_publisert,annonse_kilde&oppdragsnr=in.({",".join(nrs)})'
    for row in json.loads(http(q, SBH)): liv[str(row['oppdragsnr'])] = row
for r in signed_rows:
    L = liv.get(str(r['nr']))
    pub = L.get('annonse_publisert') if L else None
    r['publisert'] = pub; r['publisert_kilde'] = L.get('annonse_kilde') if L else None
    if pub:
        dager = (dt.date.fromisoformat(pub[:10]) - dt.date.fromisoformat(r['signert'])).days
        r['dager_til_publisert'] = dager
        add(isoweek(dt.date.fromisoformat(r['signert'])), r['megler'], 'publisert_7d', 1 if dager <= 7 else 0)
        add(isoweek(dt.date.fromisoformat(r['signert'])), r['megler'], 'publisert_kjent', 1)
NOTES.append(f'Publisert: {sum(1 for r in signed_rows if r.get("publisert"))} av {len(signed_rows)} signerte har publiseringsdato i oppdrag_livslop (resten = ikke backfillet ennå)')

# ── Utskrift ────────────────────────────────────────────────────────────────
COLS = ['kontakter', 'kontakter_fsbo', 'leads', 'tilbud_sendt', 'signert', 'publisert_7d', 'publisert_kjent', 'solgt', 'provisjon']
weeks = sorted(W.keys())
print('\n'.join('· ' + n for n in NOTES)); print()
hdrline = f'{"uke":8} {"megler":8} ' + ' '.join(f'{c:>15}' for c in COLS)
print(hdrline); print('-' * len(hdrline))
for wk in weeks:
    for who in ['Sindre', 'Henrik', 'Daniel', 'Philip', 'ukjent']:
        row = W[wk].get(who)
        if not row: continue
        print(f'{wk:8} {who:8} ' + ' '.join(f'{int(row.get(c, 0)) if c != "provisjon" else round(row.get(c, 0)):>15,}' for c in COLS))
    tot = {c: sum(W[wk][w].get(c, 0) for w in W[wk]) for c in COLS}
    print(f'{wk:8} {"SUM":8} ' + ' '.join(f'{int(tot[c]) if c != "provisjon" else round(tot[c]):>15,}' for c in COLS)); print()
print('Signerte oppdragsavtaler (Oneflow) i vinduet:')
for r in sorted(signed_rows, key=lambda x: x['signert']):
    print(f"  {r['signert']}  {r['megler']:7} {r['nr'] or '—':6} {r['navn'][:60]:60} publisert={r.get('publisert') or '?'} ({r.get('publisert_kilde') or '-'})")
if OUT:
    json.dump({'generert': TODAY.isoformat(), 'fra': FRA.isoformat(), 'notater': NOTES,
               'uker': {wk: {who: dict(W[wk][who]) for who in W[wk]} for wk in weeks}, 'signerte': signed_rows}, open(OUT, 'w'), ensure_ascii=False, indent=1)
    print('\nskrevet', OUT)
